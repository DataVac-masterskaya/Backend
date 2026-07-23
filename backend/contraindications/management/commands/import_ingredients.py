import re
from django.core.management.base import BaseCommand, CommandError

from django.db import transaction
from openpyxl import load_workbook

from reference_books.models import Ingredients


INGREDIENT_BLOCKS_PATTERN = re.compile(
    r'(?P<title>'
    r'Действующ\w*\s+веществ\w*'
    r'|Вспомогательн\w*\s+веществ\w*'
    r'|Адъювант\w*'
    r')\s*:\s*'
    r'(?P<content>.*?)(?='
    r'\s*(?:Действующ\w*\s+веществ\w*'
    r'|Вспомогательн\w*\s+веществ\w*'
    r'|Адъювант\w*)\s*:|$)',
    flags=re.IGNORECASE | re.DOTALL,
)

ACTIVE_START_MARKERS = [
    'доза содержит',
    'одна доза содержит',
    'каждая доза содержит',
    '1 доза'
]

def extract_ingredient_blocks(text):
    blocks = {}

    for match in INGREDIENT_BLOCKS_PATTERN.finditer(text):
        title = match.group('title').lower()
        content = match.group('content').strip()

        if 'действ' in title:
            key = 'active'
        elif 'вспом' in title:
            key = 'auxiliary'
        elif 'адъювант' in title:
            key = 'adjuvant'
        else:
            continue

        blocks[key] = content

    return blocks

def clean_ingredient_text(ingredients):
    ingredients = ingredients.replace('\xa0', ' ')
    ingredients = ' '.join(ingredients.split())
    ingredients = ingredients.strip(' .;,:-')
    return ingredients

def split_ingredients(text):
    ingredients = []
    current = []
    brackets = 0

    for i, char in enumerate(text):
        if char == '(':
            brackets += 1
        elif char == ')':
            brackets -= 1

        is_decimal = (
            char == ','
            and i > 0
            and i < len(text) - 1
            and text[i - 1].isdigit()
            and text[i + 1].isdigit()
        )

        is_separator = (
            char == '•'
            or (char == ',' and not is_decimal)
        )

        if is_separator and brackets == 0:
            ingredient = ''.join(current).strip(' .;:-')
            if ingredient:
                ingredients.append(ingredient)
            current = []
        else:
            current.append(char)

    ingredient = ''.join(current).strip(' .;:-')
    if ingredient:
        ingredients.append(ingredient)

    return ingredients

class Command(BaseCommand):
    help = 'Импортирует ингредиенты'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Путь к excel-файлу')

    def handle(self, *args, **options):
        ingredient_count = 0
        ingredients_set = set()
        ingredients_list = list()
        file_path = options['excel_file']
        try:
            wb = load_workbook(file_path, read_only=True)
            ws = wb['ingredients_text']
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            vaccines_id = headers.index('vaccine_id')
            ingredients_text = headers.index('ingredients_text_w/o_html')
            try:
                with transaction.atomic():
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        vaccine_id = int(row[vaccines_id])
                        ingredients_str = row[ingredients_text]

                        if vaccine_id is None:
                            raise CommandError('В файле обнаружена строка без vaccine_id.')

                        if ingredients_str is None:
                            continue

                        ingredients_blocks = extract_ingredient_blocks(ingredients_str)

                        for ingredient_type, block in ingredients_blocks.items():
                            block = clean_ingredient_text(block)

                            for ingredient in split_ingredients(block):
                                ingredient_count += 1
                                ingredients_set.add(ingredient)
                                ingredients_list.append(ingredient)
                                # Ingredients.objects.get_or_create(
                                #     name=ingredient,
                                #     type=ingredient_type,
                                # )
                                self.stdout.write(
                                    f' {vaccine_id} - {ingredient_type}: {ingredient}')
            finally:
                ingredients_list.sort()
                print(ingredient_count)
                print(len(ingredients_set))
                print(ingredients_list)
                wb.close()
        except Exception as e:
            raise CommandError(str(e))
