from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from contraindications.models import Contraindication


class Command(BaseCommand):
    help = 'Импортирует противопоказания из указанного excel-файла со страницы contraindications_list'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Путь к excel-файлу')

    def handle(self, *args, **options):
        file_path = options['excel_file']
        try:
            wb = load_workbook(file_path, read_only=True)
            ws = wb['contraindications_list']
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            name_idx = headers.index('contraindication_name')
            id_idx = headers.index('contraindication_ID')
        except Exception as e:
            raise CommandError(f'Ошибка при чтении файла: {e}')
        count_add = 0
        seen_ids = set()
        try:
            with transaction.atomic():
                for row in ws.iter_rows(min_row=2, values_only=True):
                    contraindication_name = row[name_idx]
                    old_id = row[id_idx]
                    if old_id is None:
                        raise CommandError('В файле обнаружена строка без contraindication_ID.')

                    if contraindication_name is None:
                        raise CommandError(f'Для contraindication_ID={old_id} отсутствует название.')
                    if old_id in seen_ids:
                        raise CommandError(
                            'В файле обнаружен элемент с повторяющимся '
                            f'contraindication_ID: {old_id}, '
                            f'contraindication_name: {contraindication_name}'
                        )
                    seen_ids.add(old_id)
                    _, created = Contraindication.objects.update_or_create(
                        old_id=old_id,
                        defaults={
                            'name': contraindication_name,
                        },
                    )
                    if created:
                        count_add += 1

            self.stdout.write(f'Импорт противопоказаний завершён. Добавлено: {count_add}. ')
        finally:
            wb.close()
