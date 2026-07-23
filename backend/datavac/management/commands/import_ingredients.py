import re
from django.core.management.base import BaseCommand, CommandError

from django.db import transaction
from openpyxl import load_workbook

from reference_books.models import Ingredients


INGREDIENT_BLOCKS_PATTERN = re.compile(
    r'(?P<title>'
    r'Действующ(?:ее|ие)\s+веществ(?:о|а)'
    r'|Вспомогательн(?:ое|ые)\s+веществ(?:о|а)'
    r'|Адъювант(?:ы)?'
    r')\s*:\s*'
    r'(?P<content>.*?)(?='
    r'\n\s*(?:Действующ(?:ее|ие)\s+веществ(?:о|а)'
    r'|Вспомогательн(?:ое|ые)\s+веществ(?:о|а)'
    r'|Адъювант(?:ы)?'
    r')\s*:|$)',
    flags=re.IGNORECASE | re.DOTALL,
)


def extract_ingredient_blocks(text):
    blocks = {}

    for match in INGREDIENT_BLOCKS_PATTERN.finditer(text):
        title = match.group('title').lower()
        content = match.group('content').strip()

        if 'действ' in title:
            key = 'active'
        elif 'вспом' in title:
            key = 'auxiliary'
        elif 'адъювант' in title:
            key = 'adjuvant'
        else:
            continue

        blocks[key] = content

    return blocks

def clean_ingredient_text(ingredients):
    ingredients = ingredients.replace('\xa0', ' ')
    ingredients = ' '.join(ingredients.split())
    ingredients = ingredients.strip(' .;,:-')
    return ingredients

def split_ingredients(text):
    ingredients = []
    current = []
    brackets = 0

    for char in text:
        if char == '(':
            brackets += 1
        elif char == ')':
            brackets -= 1

        if char == ',' and brackets == 0:
            ingredient = ''.join(current).strip()
            if ingredient:
                ingredients.append(ingredient)
            current = []
        else:
            current.append(char)

    ingredient = ''.join(current).strip()
    if ingredient:
        ingredients.append(ingredient)

    return ingredients


class Command(BaseCommand):
    help = 'Импортирует ингредиенты'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Путь к excel-файлу')

    def handle(self, *args, **options):
        file_path = options['excel_file']
        try:
            wb = load_workbook(file_path, read_only=True)
            ws = wb['ingredients_text']
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            vaccines_id = headers.index('vaccine_id')
            ingredients_text = headers.index('ingredients_text_w/o_html')
            try:
                with transaction.atomic():
                for row in ws.iter_rows(min_row=2, values_only=True):
                    vaccine_id = row[vaccines_id]
                    ingredients_str = row[ingredients_text]

                    if vaccine_id is None:
                        raise CommandError('В файле обнаружена строка без vaccine_id.')

                    if ingredients_str is None:
                        continue

                    ingredients_blocks = extract_ingredient_blocks(ingredients_str)

                    for ingredient_type, block in ingredients_blocks.items():
                        block = clean_ingredient_text(block)

                        for ingredient in split_ingredients(block):
                            # Ingredients.objects.get_or_create(
                            #     name=ingredient,
                            #     type=ingredient_type,
                            # )
                            self.stdout.write(
                                f'{ingredient_type}: {ingredient}'
)
        finally:
            wb.close()